"""Execute the committed Excel plan in row order and write auditable results."""
import json
import os
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
PLAN = ROOT / "qa/Role-AI_unit_tests.xlsx"
OUT = ROOT / "qa/results"
OUT.mkdir(parents=True, exist_ok=True)
workbook = load_workbook(PLAN)
sheet = workbook["단위테스트"]
cases = list(sheet.iter_rows(min_row=2, values_only=True))
selectors = [row[9] for row in cases]
if len(set(selectors)) != len(selectors):
    raise SystemExit("Duplicate selectors in Excel")
if [int(row[0]) for row in cases] != list(range(1, len(cases) + 1)):
    raise SystemExit("Excel order must be contiguous")


class OrderedPlan:
    def __init__(self):
        self.collected = []
        self.started = []
        self.reports = {}

    def pytest_collection_modifyitems(self, session, config, items):
        def selector(item):
            return item.nodeid.split("[", 1)[0]
        found = {selector(item) for item in items}
        if found != set(selectors):
            raise pytest.UsageError(
                f"Excel/code mismatch: missing={set(selectors) - found}, extra={found - set(selectors)}"
            )
        rank = {node: i for i, node in enumerate(selectors)}
        items.sort(key=lambda item: rank[selector(item)])
        self.collected = [item.nodeid for item in items]

    def pytest_runtest_logstart(self, nodeid, location):
        self.started.append(nodeid)

    def pytest_runtest_logreport(self, report):
        self.reports.setdefault(report.nodeid, []).append({
            "phase": report.when, "outcome": report.outcome,
            "seconds": report.duration,
            "detail": str(report.longrepr) if report.failed or report.skipped else "",
        })


plugin = OrderedPlan()
exit_code = int(pytest.main(
    ["tests", "-q", "--junitxml=" + str(OUT / "test-results.xml")], plugins=[plugin]
))
ordered = plugin.started == plugin.collected and bool(plugin.collected)
results = []
for row in cases:
    nodes = [node for node in plugin.collected if node.split("[", 1)[0] == row[9]]
    reports = [report for node in nodes for report in plugin.reports.get(node, [])]
    completed = all(any(r["phase"] == "call" for r in plugin.reports.get(node, []))
                    for node in nodes)
    passed = bool(nodes) and completed and all(r["outcome"] == "passed" for r in reports)
    status = "PASS" if passed and ordered else (
        "FAIL" if any(r["outcome"] == "failed" for r in reports) else "미완료"
    )
    result = {
        "id": row[1], "selector": row[9], "status": status, "executions": len(nodes),
        "seconds": round(sum(r["seconds"] for r in reports), 4),
        "detail": "\n".join(r["detail"] for r in reports if r["detail"]),
    }
    results.append(result)
    index = int(row[0]) + 1
    for column, value in enumerate(
        [status, len(nodes), result["seconds"], result["detail"][:30000]], start=11
    ):
        sheet.cell(index, column, value)

summary = {
    "commit": os.getenv("GITHUB_SHA", ""),
    "run_url": "https://github.com/" + os.getenv("GITHUB_REPOSITORY", "wjdwodhks35/Role-AI")
               + "/actions/runs/" + os.getenv("GITHUB_RUN_ID", ""),
    "utc": datetime.now(timezone.utc).isoformat(),
    "plan": str(PLAN.relative_to(ROOT)),
    "ordered": ordered, "exit_code": exit_code,
    "scenarios": len(cases), "executions": len(plugin.started),
    "passed_scenarios": sum(r["status"] == "PASS" for r in results),
    "results": results,
    "execution_order": plugin.started,
}
defects_path = ROOT / "qa/defects.json"
if defects_path.exists():
    defects = json.loads(defects_path.read_text(encoding="utf-8"))
    defect_sheet = workbook["결함이력"]
    defect_sheet.delete_rows(2, defect_sheet.max_row)
    for defect in defects:
        case_result = next(r for r in results if r["id"] == defect["case"])
        defect_sheet.append([
            defect["id"], defect["case"], defect["before"] + " / " + defect["evidence"],
            defect["fix"], case_result["status"] + " / " + summary["run_url"],
        ])
for key in ["commit", "run_url", "utc", "ordered", "exit_code", "scenarios",
            "executions", "passed_scenarios"]:
    workbook["요약"].append([key, str(summary[key])])
details = workbook.create_sheet("실행상세")
details.append(["실행순서", "pytest node", "단계", "결과", "소요초", "실패상세"])
for index, node in enumerate(plugin.started, 1):
    for report in plugin.reports.get(node, []):
        details.append([index, node, report["phase"], report["outcome"],
                        report["seconds"], report["detail"][:30000]])
for ws in workbook:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23395B")
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = (
            12 if column[0].column <= 2 else 42
        )
workbook.save(OUT / "Role-AI_unit_tests_results.xlsx")
(OUT / "results.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
)
print("QA_RESULTS_JSON=" + json.dumps(summary, ensure_ascii=False, separators=(",", ":")))
raise SystemExit(exit_code or (0 if ordered and all(
    r["status"] == "PASS" for r in results
) else 1))

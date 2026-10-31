"""Append only the new foundation test selectors to the existing Excel plan."""

import ast
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def main():
    path = ROOT / "qa/Role-AI_unit_tests.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["단위테스트"]
    existing = {row[9] for row in sheet.iter_rows(min_row=2, values_only=True)}
    added = 0
    for filename, requirement, category in [
        ("test_code_execution_permissions.py", "REQ-CODE-EXECUTION", "umask·비root bind·호스트 부모 비공개"),
        ("test_agent_registry.py", "REQ-AGENT-CONFIG", "역할 구성·권한 경계"),
        ("test_baseline_manifest.py", "REQ-BASELINE", "읽기 전용 기준선"),
        ("test_coding_workspace.py", "REQ-CODING-WORKSPACE", "코딩 작업·격리 파일 변경"),
        ("test_strategy_events.py", "REQ-STRATEGY-RECORDER", "최소 전략 경험 기록"),
        ("test_coding_runtime.py", "REQ-CODING-RUNTIME", "코딩 실행·복구·체크포인트"),
        ("test_coding_stability.py", "REQ-CODING-STABILITY", "코딩 모델 연동·패치 복구·토큰 계측"),
        ("test_project_scopes.py", "REQ-PROJECT-SCOPE", "프로젝트·버전 공유·파생 출처"),
        ("test_scope_migrations.py", "REQ-SCOPE-MIGRATION", "추가형 범위 이행·복제본·복구"),
        ("test_scope_permissions.py", "REQ-SCOPE-PERMISSION", "인증 세션·계약·자료 권한 Gate"),
        ("test_scoped_reads.py", "REQ-SCOPED-DATA", "현재 범위 읽기 Adapter"),
        ("test_scope_data_paths.py", "REQ-SCOPED-DATA", "문서·기억·대화·보관 경로"),
        ("test_task_runtime_endpoint.py", "REQ-BENCHMARK-ENDPOINT", "명시적 사용자 Unix socket·실시간 바인딩·자원 preflight"),
        ("test_task_comparison.py", "REQ-TASK-COMPARISON", "동일10과제·원문/조건 고정·비교/미측정·기준선"),
        ("test_task_metrics.py", "REQ-TASK-METRICS", "실측/미측정·실패 분모·GPU 샘플·증거 보존"),
        ("test_task_delegation.py", "REQ-TASK-DELEGATION", "공통 계약·기존 Coding Agent·독립 채점 연결"),
        ("test_windows_loopback_transport.py", "REQ-WINDOWS-LOOPBACK", "고정 Windows loopback·무권한 데이터 전달"),
        ("test_independent_evaluation.py", "REQ-INDEPENDENT-GRADING", "혼합10과제·공개/숨김·외부 oracle·증거 보존"),
        ("test_evaluation_reuse.py", "REQ-EVALUATION-REUSE", "검증된 PR30 기반·PR31 오류/체크포인트 선별"),
        ("test_phase_one_integration.py", "REQ-PHASE-ONE-INTEGRATION", "공통 계약·격리·승인·복구 통합"),
        ("test_task_resources.py", "REQ-TASK-RESOURCES", "제한 재시도·예약 ledger·시간/GPU·복구"),
        ("test_task_checkpoints.py", "REQ-TASK-STATE", "상태 전이·체크포인트·중복 실행·재개"),
        ("test_task_approvals.py", "REQ-TASK-APPROVAL", "유효 승인·도구 범위·계약 바인딩"),
        ("test_task_approval_migrations.py", "REQ-TASK-APPROVAL", "v6 승인 기록·보존·복구"),
        ("test_scope_data_migrations.py", "REQ-SCOPED-DATA", "v5 추가형 이행·원문 보존"),
    ]:
        tree = ast.parse((ROOT / "tests" / filename).read_text())
        for node in tree.body:
            if not isinstance(node, ast.FunctionDef) or not node.name.startswith("test_"):
                continue
            selector = f"tests/{filename}::{node.name}"
            if selector in existing:
                continue
            order = sheet.max_row
            previous = sheet.max_row
            sheet.append(
                [
                    str(order),
                    f"UT-{order:03d}",
                    requirement,
                    category,
                    "단위/API",
                    node.name,
                    "합성 설정·임시 DB·로컬 fixture",
                    "설정·인증·격리·무변경 조건을 pytest로 검사",
                    "무권한 접근·실행 차단, 비밀 및 기존 상태 보존",
                    selector,
                    "미실행",
                    None,
                    None,
                    None,
                ]
            )
            for target, source in zip(sheet[sheet.max_row], sheet[previous]):
                target._style = copy(source._style)
                target.alignment = copy(source.alignment)
            existing.add(selector)
            added += 1
    changed = bool(added)
    for row in workbook["요약"]:
        if row[0].value == "계획":
            description = f"{sheet.max_row - 1}개 시나리오; 매개변수별 실행 수는 결과에 기록"
            if row[1].value != description:
                row[1].value = description
                changed = True
    coverage = workbook["요구사항커버리지"]
    names = {row[0].value for row in coverage}
    for name, scope in [
        (
            "REQ-AGENT-CONFIG",
            "여섯 역할 정의/설정 검증/권한 확대 차단/로컬 소유자 조회/기존 대화 유지",
        ),
        ("REQ-BASELINE", "소스·설정 해시/Git 식별/원문·비밀 제외/기존 보고서 덮어쓰기 방지"),
        (
            "REQ-CODING-WORKSPACE",
            "작업·변경 승인 해시/격리 worktree/원본 보존/경로·수정량 제한/실행 금지/diff",
        ),
        ("REQ-STRATEGY-RECORDER", "최소 메타데이터/범위/장애/비밀 미저장"),
        ("REQ-CODING-RUNTIME", "탐색/반복 수정/격리 명령/오류 복구/체크포인트/중단·재개/예산·승인"),
        (
            "REQ-CODING-STABILITY",
            "thinking 제어/응답 오류 분류/작업 메모리·반복 중단/패치 정규화·복구/실측 토큰 정산/keep_alive",
        ),
        ("REQ-PROJECT-SCOPE", "개인 보관함/소유권/정확 버전 읽기 Grant/철회/파생 기억 출처 보존"),
        ("REQ-SCOPE-MIGRATION", "합성 v3→v4/원문 보존/재실행/실패 rollback/복제 DB·백업 복원/운영 미적용"),
        ("REQ-SCOPE-PERMISSION", "인증 세션/소유권/정확 버전/정책 세대/철회/파생 출처 읽기 Gate, 실행 권한 확대 금지"),
        ("REQ-BENCHMARK-ENDPOINT", "A/B 로컬 Unix socket 선택·소유권/링크/교체 거부·Rootless 자원 제한·기존 기본값 보존"),
        ("REQ-TASK-COMPARISON", "공개 기준선/독립 채점/버전·seed·조건/실제·합성/회귀 차단·불충분 보류"),
        ("REQ-TASK-METRICS", "실제/합성 구분·토큰 null·GPU 관측·미채점 분모·실행 중복 보존"),
        ("REQ-TASK-DELEGATION", "기존 Coding Agent/common gate·예산 ledger/공개·독립 채점 분리·취소/철회 차단"),
        ("REQ-WINDOWS-LOOPBACK", "고정 localhost11434/상수 PowerShell·stdin 데이터/설치·원격·generate 차단"),
        ("REQ-INDEPENDENT-GRADING", "합성 Role-AI5/소형5·공개2/숨김4/정답 외부 채점·위조/철회/cleanup 차단·로컬 보고서"),
        ("REQ-EVALUATION-REUSE", "PR30 현재 기반 재사용/PR31 typed checkpoint 오류·bounded type-only 분류/저장 보존/HostRunner·자동테스트 미도입"),
        ("REQ-PHASE-ONE-INTEGRATION", "P1-01~07/승인 우회·철회·소유권/원본 보존/DB+공통 checkpoint 비활성 복원"),
        ("REQ-TASK-RESOURCES", "실패 분류·제한 재시도/단계·시간·토큰 예약/GPU admission/미측정 분리/중단 복구·무재실행"),
        ("REQ-TASK-STATE", "불변 계약·실행 상태 분리/원자 체크포인트/소유권·CAS·lock/중단·재개/독립 검증"),
        ("REQ-TASK-APPROVAL", "인증·CSRF/계약 digest/세션·정책·도구 설정 바인딩/철회/경로 보호/v6 추가형 이행·복구"),
        ("REQ-SCOPED-DATA", "개인 기본 격리/명시 버전 공유/보관/대화 고정/출처 추적/v5 보존·복구"),
    ]:
        if name not in names:
            coverage.append([name, scope, "자동 검증; 실제 모델·학습·운영 반영 없음"])
            changed = True
    if changed:
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(path)
    print(f"Added {added} foundation scenarios; existing plan rows preserved.")


if __name__ == "__main__":
    main()
